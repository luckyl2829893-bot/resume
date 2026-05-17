import io
import re
from fpdf import FPDF
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def clean_pdf_text(text: str) -> str:
    """
    Cleans Unicode characters that FPDF Helvetica doesn't support by replacing them
    with standard ASCII or Latin-1 equivalents to prevent encoding crashes.
    """
    replacements = {
        "\u201c": '"',  # Left double quote
        "\u201d": '"',  # Right double quote
        "\u2018": "'",  # Left single quote
        "\u2019": "'",  # Right single quote
        "\u2013": "-",  # En dash
        "\u2014": "-",  # Em dash
        "\u2022": chr(183),  # Bullet point symbol (replaced by safe Latin-1 middle dot)
        "\u2026": "...",# Ellipsis
        "\u00a0": " ",  # Non-breaking space
        "\u20ac": "EUR",# Euro symbol
        "\u20b9": "INR",# Rupee symbol
        "\u00ae": "(R)",# Registered symbol
        "\u00a9": "(C)",# Copyright symbol
        "\u2122": "(TM)",# Trademark symbol
    }
    for orig, rep in replacements.items():
        text = text.replace(orig, rep)
        
    # Rebuild string character by character, keeping only Latin-1 compatible characters
    cleaned_chars = []
    for char in text:
        try:
            char.encode("latin-1")
            cleaned_chars.append(char)
        except UnicodeEncodeError:
            cleaned_chars.append("?")
    return "".join(cleaned_chars)

def write_formatted_text(pdf, text: str, font_family: str, size: float, line_height: float = 4.5):
    """
    Splits text by '**' and writes alternating normal vs bold chunks to the PDF flow.
    """
    parts = text.split("**")
    for idx, part in enumerate(parts):
        if idx % 2 == 1:
            pdf.set_font(font_family, "B", size)
        else:
            pdf.set_font(font_family, "", size)
        pdf.write(line_height, part)

def parse_tailored_sections(tailored_text: str) -> dict:
    """
    Parses tailored resume text into a structured dictionary.
    Handles structured NAME/CONTACT/SECTION layout and falls back gracefully to
    unstructured line parsing if tags are missing.
    """
    result = {
        "name": "Candidate Name",
        "contact": "",
        "sections": {}
    }
    
    text = tailored_text.strip()
    if text.startswith("```"):
        text = text[text.find("\n")+1:]
    if text.endswith("```"):
        text = text[:text.rfind("```")]
    text = text.strip()
    
    lines = text.split("\n")
    current_section = None
    section_lines = []
    
    # 1. Standard parser for structured NAME/CONTACT/---SECTION:
    has_structure = any("---SECTION:" in line for line in lines)
    
    if has_structure:
        for line in lines:
            line_strip = line.strip()
            if not line_strip:
                continue
            if line_strip.startswith("NAME:"):
                result["name"] = line_strip[len("NAME:"):].strip()
            elif line_strip.startswith("CONTACT:"):
                result["contact"] = line_strip[len("CONTACT:"):].strip()
            elif line_strip.startswith("---SECTION:") and line_strip.endswith("---"):
                if current_section:
                    result["sections"][current_section] = "\n".join(section_lines).strip()
                    section_lines = []
                current_section = line_strip[len("---SECTION:"): -3].strip().upper()
            else:
                if current_section:
                    section_lines.append(line)
        if current_section and section_lines:
            result["sections"][current_section] = "\n".join(section_lines).strip()
            
    # 2. Fallback parser for unstructured text
    else:
        KNOWN_HEADERS = [
            "SUMMARY", "PROFILE SUMMARY", "EXPERIENCE", "PROFESSIONAL EXPERIENCE", 
            "EDUCATION", "SKILLS", "TECHNICAL SKILLS", "PROJECTS", "INTERESTS", 
            "PROFESSIONAL SKILLS", "LINKS", "INTEREST", "PUBLICATIONS"
        ]
        
        current_section = "SUMMARY"
        first_line = True
        
        for line in lines:
            line_strip = line.strip()
            if not line_strip:
                continue
            if first_line:
                result["name"] = line_strip.replace("#", "").replace("**", "").strip()
                first_line = False
                continue
            if not result["contact"] and any(term in line_strip.lower() for term in ["|", "linkedin", "github", "email", "@"]):
                result["contact"] = line_strip.replace("**", "").replace("*", "").strip()
                continue
                
            clean_line = line_strip.replace("#", "").replace("**", "").strip().upper()
            if clean_line in KNOWN_HEADERS or (clean_line.isupper() and len(clean_line) < 30):
                if current_section:
                    result["sections"][current_section] = "\n".join(section_lines).strip()
                current_section = clean_line
                section_lines = []
            else:
                section_lines.append(line)
                
        if current_section and section_lines:
            result["sections"][current_section] = "\n".join(section_lines).strip()
            
    return result

def build_pdf(tailored_text: str, filename: str = "resume.pdf", layout_profile: dict = None) -> bytes:
    """
    Generates a high-fidelity, single-column, ATS-safe PDF from tailored resume text.
    Strictly replicates the visual layout fingerprint (fonts, colors, structure, sections)
    extracted from the original user's resume.
    """
    if not layout_profile:
        # Default elegant Deep Teal profile fallback
        layout_profile = {
            "accent_color": (19, 148, 148),
            "accent_color_hex": "139494",
            "name_font_size": 22,
            "name_align": "C",
            "section_order": ["SUMMARY", "EXPERIENCE", "PROJECTS", "EDUCATION", "SKILLS", "INTERESTS"],
            "bullet_char": chr(183),
            "heading_font_size": 11,
            "body_font_size": 9.5
        }
        
    accent_r, accent_g, accent_b = layout_profile.get("accent_color", (19, 148, 148))
    bullet_char = layout_profile.get("bullet_char", "•")
    
    # Ensure PDF bullet char is encodable in Latin-1
    try:
        bullet_char.encode("latin-1")
    except UnicodeEncodeError:
        bullet_char = chr(183)  # Middle dot fallback
        
    parsed = parse_tailored_sections(tailored_text)
    
    pdf = FPDF()
    pdf.add_page()
    # A4 margins: 15mm left/right, 12mm top/bottom
    pdf.set_margins(15, 12, 15)
    pdf.set_auto_page_break(auto=True, margin=12)
    
    # Header Name
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Helvetica", "B", layout_profile.get("name_font_size", 22))
    name_align = layout_profile.get("name_align", "C")
    pdf.cell(0, 8, clean_pdf_text(parsed["name"]), align=name_align, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1)
    
    # Contact details row
    if parsed["contact"]:
        pdf.set_text_color(80, 80, 80)
        pdf.set_font("Helvetica", "", 9.5)
        # Standard pipe-separated formatting for clean visual layout
        contact_text = parsed["contact"].replace("**", "").replace("*", "").strip()
        pdf.cell(0, 5, clean_pdf_text(contact_text), align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(4)
        
    # Render sections according to the exact extracted section order
    for sec_name in layout_profile.get("section_order", []):
        parsed_sec_key = None
        for key in parsed["sections"].keys():
            if sec_name.upper() in key or key in sec_name.upper():
                parsed_sec_key = key
                break
                
        if not parsed_sec_key:
            continue
            
        content = parsed["sections"][parsed_sec_key]
        if not content:
            continue
            
        # Draw Section Header
        pdf.ln(3)
        pdf.set_text_color(accent_r, accent_g, accent_b)
        pdf.set_font("Helvetica", "B", layout_profile.get("heading_font_size", 11))
        pdf.cell(0, 6, sec_name.upper(), new_x="LMARGIN", new_y="NEXT")
        
        # Horizontal accent rule
        pdf.set_draw_color(accent_r, accent_g, accent_b)
        pdf.set_line_width(0.35)
        pdf.line(15, pdf.get_y() + 0.5, 195, pdf.get_y() + 0.5)
        pdf.ln(2.5)
        
        # Render Section Content Body
        pdf.set_text_color(30, 30, 30)
        lines = content.split("\n")
        
        for line in lines:
            line_clean = clean_pdf_text(line.strip())
            if not line_clean:
                pdf.ln(1.5)
                continue
                
            # Detect sub-headings (e.g., project titles starting with ##PROJECT:)
            if line_clean.startswith("##PROJECT:"):
                proj_title = line_clean[len("##PROJECT:"):].strip()
                pdf.set_font("Helvetica", "B", layout_profile.get("body_font_size", 10))
                pdf.set_text_color(0, 0, 0)
                pdf.set_x(15)
                pdf.write(4.5, proj_title)
                pdf.ln(4.5)
                pdf.set_text_color(30, 30, 30)
                continue
                
            # Detect bullet points
            is_bullet = line_clean.startswith("*") or line_clean.startswith("-") or line_clean.startswith("•") or line_clean.startswith(chr(183))
            
            if is_bullet:
                bullet_text = re.sub(r'^[\*\-\•' + chr(183) + r']\s*', '', line_clean).strip()
                pdf.set_font("Helvetica", "B", layout_profile.get("body_font_size", 9.5))
                pdf.set_x(15)
                pdf.write(4.5, bullet_char + " ")
                
                # Shift left margin dynamically to wrap bullet points with clean indentation
                pdf.set_left_margin(20)
                pdf.set_x(20)
                write_formatted_text(pdf, bullet_text, "Helvetica", layout_profile.get("body_font_size", 9.5), 4.5)
                pdf.ln(4.5)
                pdf.set_left_margin(15)
            else:
                pdf.set_x(15)
                write_formatted_text(pdf, line_clean, "Helvetica", layout_profile.get("body_font_size", 9.5), 4.5)
                pdf.ln(4.5)
                
        pdf.ln(1.5)
        
    return bytes(pdf.output())

def build_docx(tailored_text: str, filename: str = "resume.docx", layout_profile: dict = None) -> bytes:
    """
    Generates a clean, ATS-safe, layout-mirrored DOCX from tailored resume text.
    Uses Calibri and mirrors bottom heading borders using custom paragraph XML injection.
    """
    if not layout_profile:
        layout_profile = {
            "accent_color": (19, 148, 148),
            "accent_color_hex": "139494",
            "name_font_size": 22,
            "name_align": "C",
            "section_order": ["SUMMARY", "EXPERIENCE", "PROJECTS", "EDUCATION", "SKILLS", "INTERESTS"],
            "bullet_char": "•",
            "heading_font_size": 11,
            "body_font_size": 10
        }
        
    accent_rgb = RGBColor(*layout_profile.get("accent_color", (19, 148, 148)))
    hex_color = layout_profile.get("accent_color_hex", "139494")
    
    parsed = parse_tailored_sections(tailored_text)
    doc = Document()
    
    # 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Header Name
    p_name = doc.add_paragraph()
    if layout_profile.get("name_align", "C") == "C":
        p_name.alignment = 1  # Center
    else:
        p_name.alignment = 0  # Left
    p_name.paragraph_format.space_after = Pt(2)
    run_name = p_name.add_run(parsed["name"])
    run_name.bold = True
    run_name.font.size = Pt(20)
    run_name.font.name = "Calibri"
    run_name.font.color.rgb = RGBColor(0, 0, 0)
    
    # Contact row
    if parsed["contact"]:
        p_contact = doc.add_paragraph()
        p_contact.alignment = 1  # Center
        p_contact.paragraph_format.space_after = Pt(6)
        contact_text = parsed["contact"].replace("**", "").replace("*", "").strip()
        run_contact = p_contact.add_run(contact_text)
        run_contact.font.size = Pt(9.5)
        run_contact.font.name = "Calibri"
        run_contact.font.color.rgb = RGBColor(80, 80, 80)
        
    for sec_name in layout_profile.get("section_order", []):
        parsed_sec_key = None
        for key in parsed["sections"].keys():
            if sec_name.upper() in key or key in sec_name.upper():
                parsed_sec_key = key
                break
                
        if not parsed_sec_key:
            continue
            
        content = parsed["sections"][parsed_sec_key]
        if not content:
            continue
            
        # Section Header Paragraph
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(12)
        p_head.paragraph_format.space_after = Pt(3)
        p_head.paragraph_format.keep_with_next = True
        
        run_head = p_head.add_run(sec_name.upper())
        run_head.bold = True
        run_head.font.size = Pt(layout_profile.get("heading_font_size", 11))
        run_head.font.name = "Calibri"
        run_head.font.color.rgb = accent_rgb
        
        # Inject paragraph bottom border to replicate horizontal accent line divider
        pPr = p_head._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), hex_color)
        pBdr.append(bottom)
        pPr.append(pBdr)
        
        # Render Section Content lines
        lines = content.split("\n")
        for line in lines:
            line_clean = line.strip()
            if not line_clean:
                continue
                
            if line_clean.startswith("##PROJECT:"):
                proj_title = line_clean[len("##PROJECT:"):].strip()
                p_proj = doc.add_paragraph()
                p_proj.paragraph_format.space_before = Pt(4)
                p_proj.paragraph_format.space_after = Pt(1)
                run_p = p_proj.add_run(proj_title)
                run_p.bold = True
                run_p.font.name = "Calibri"
                run_p.font.size = Pt(10)
                run_p.font.color.rgb = RGBColor(0, 0, 0)
                continue
                
            is_bullet = line_clean.startswith("*") or line_clean.startswith("-") or line_clean.startswith("•")
            
            if is_bullet:
                bullet_text = re.sub(r'^[\*\-\•]\s*', '', line_clean).strip()
                p_item = doc.add_paragraph(style='List Bullet')
                p_item.paragraph_format.space_before = Pt(0)
                p_item.paragraph_format.space_after = Pt(2)
                p_item.paragraph_format.left_indent = Inches(0.25)
                text_to_parse = bullet_text
            else:
                p_item = doc.add_paragraph()
                p_item.paragraph_format.space_before = Pt(0)
                p_item.paragraph_format.space_after = Pt(3)
                text_to_parse = line_clean
                
            # Parse inline bold blocks
            parts = text_to_parse.split("**")
            for idx, part in enumerate(parts):
                run = p_item.add_run(part)
                run.font.name = "Calibri"
                run.font.size = Pt(10)
                if idx % 2 == 1:
                    run.bold = True
                else:
                    run.bold = False
                    
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
